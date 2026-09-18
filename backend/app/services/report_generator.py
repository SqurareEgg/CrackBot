"""
Real PDF and DOCX report generation for building safety inspection reports.
Uses reportlab (PDF) and python-docx (DOCX) with Korean font support.
"""
import io
import zipfile
from datetime import datetime
from pathlib import Path

from sqlalchemy.orm import Session

from app.models.photo import DefectDetection
from app.models.report import Report, ReportFormat
from app.services.grade import calculate_grade, grade_recommendation

_REPORTS_DIR = Path(__file__).parent.parent.parent / "reports"
_REPORTS_DIR.mkdir(exist_ok=True)

# Korean font (Malgun Gothic is bundled with Windows)
_FONT_PATH = Path("C:/Windows/Fonts/malgun.ttf")
_FONT_BOLD_PATH = Path("C:/Windows/Fonts/malgunbd.ttf")


# ── Data gathering ──────────────────────────────────────────────

def _gather_data(report: Report, db: Session) -> dict:
    project = report.project
    if not project:
        raise ValueError("프로젝트를 찾을 수 없습니다.")

    building = project.building
    inspector = project.inspector

    zones_data = []
    all_detections: list[DefectDetection] = []

    for zone in project.zones:
        zone_dets: list[DefectDetection] = []
        for photo in zone.photos:
            zone_dets.extend(photo.detections)
        all_detections.extend(zone_dets)

        max_crack = max(
            (d.crack_width_mm for d in zone_dets if d.crack_width_mm), default=None
        )
        grade = calculate_grade(max_crack, len(zone_dets))
        zones_data.append({
            "name": zone.zone_name,
            "photo_count": len(zone.photos),
            "defect_count": len(zone_dets),
            "max_crack_mm": max_crack,
            "grade": grade,
        })

    overall_max = max(
        (d.crack_width_mm for d in all_detections if d.crack_width_mm), default=None
    )
    overall_grade = calculate_grade(overall_max, len(all_detections))

    return {
        "report_id": report.id,
        "generated_at": datetime.utcnow(),
        "template_type": str(report.template_type.value),
        "project_title": project.title,
        "building_name": building.name if building else "—",
        "building_address": building.address if building else "—",
        "start_date": str(project.start_date),
        "end_date": str(project.end_date) if project.end_date else "미정",
        "overall_grade": overall_grade,
        "overall_defect_count": len(all_detections),
        "overall_max_crack_mm": overall_max,
        "recommendation": grade_recommendation(overall_grade, overall_max),
        "zones": zones_data,
        "inspector_name": inspector.name if inspector else "—",
    }


# ── PDF ─────────────────────────────────────────────────────────

_GRADE_HEX = {
    "A": "#059669", "B": "#0284C7",
    "C": "#D97706", "D": "#DC2626", "E": "#7C3AED",
}


def _hex(h: str):
    from reportlab.lib import colors
    return colors.HexColor(h)


def _register_korean_font() -> str:
    """Register Malgun Gothic and return the font name. Falls back to Helvetica."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if _FONT_PATH.exists():
        try:
            pdfmetrics.registerFont(TTFont("MalgunGothic", str(_FONT_PATH)))
            if _FONT_BOLD_PATH.exists():
                pdfmetrics.registerFont(TTFont("MalgunGothic-Bold", str(_FONT_BOLD_PATH)))
            else:
                pdfmetrics.registerFont(TTFont("MalgunGothic-Bold", str(_FONT_PATH)))
            from reportlab.pdfbase.pdfmetrics import registerFontFamily
            registerFontFamily("MalgunGothic", normal="MalgunGothic", bold="MalgunGothic-Bold")
            return "MalgunGothic"
        except Exception:
            pass
    return "Helvetica"


def _generate_pdf(data: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    font = _register_korean_font()
    bold = font + "-Bold" if font == "MalgunGothic" else "Helvetica-Bold"

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
        title="건축물 안전점검 보고서",
    )

    PRIMARY   = _hex("#1D9E75")
    GRAY      = _hex("#6B7280")
    DARK      = _hex("#111827")
    BODY_CLR  = _hex("#374151")
    LIGHT_BG  = _hex("#F9FAFB")
    BORDER    = _hex("#E5E7EB")
    WHITE     = colors.white

    def style(name, **kw):
        kw.setdefault('fontName', font)
        return ParagraphStyle(name, **kw)

    title_s    = style("T",  fontSize=20, textColor=DARK,    alignment=1, spaceAfter=4,  leading=26)
    sub_s      = style("S",  fontSize=10, textColor=GRAY,    alignment=1, spaceAfter=14, leading=14)
    h2_s       = style("H2", fontSize=13, textColor=DARK,    spaceBefore=18, spaceAfter=8, leading=18, fontName=bold)
    body_s     = style("B",  fontSize=10, textColor=BODY_CLR, leading=15)
    label_s    = style("L",  fontSize=8,  textColor=GRAY,    leading=12)
    grade_s    = style("G",  fontSize=18, textColor=_hex(_GRADE_HEX.get(data["overall_grade"], "#374151")),
                       alignment=1, leading=24, fontName=bold)
    rec_s      = style("R",  fontSize=10, textColor=BODY_CLR, alignment=1, leading=15)

    story = []

    # ── Header ──
    story.append(Paragraph("건축물 안전점검 보고서", title_s))
    tmpl_label = "국토안전관리원 표준 양식" if data["template_type"] == "standard" else "간이 요약 보고서"
    story.append(Paragraph(tmpl_label, sub_s))
    story.append(HRFlowable(width="100%", thickness=2, color=PRIMARY, spaceAfter=14))

    # ── Project info ──
    info = [
        ["프로젝트명", data["project_title"],    "건물명",   data["building_name"]],
        ["주소",       data["building_address"], "점검 기간", f"{data['start_date']} ~ {data['end_date']}"],
        ["점검자",     data["inspector_name"],   "생성 일시", data["generated_at"].strftime("%Y-%m-%d %H:%M")],
    ]
    col_w = [28*mm, 68*mm, 28*mm, 56*mm]
    t = Table(info, colWidths=col_w)
    t.setStyle(TableStyle([
        ("FONTNAME",    (0, 0), (-1, -1), font),
        ("FONTSIZE",    (0, 0), (-1, -1), 9),
        ("TEXTCOLOR",   (0, 0), (0, -1), GRAY),
        ("TEXTCOLOR",   (2, 0), (2, -1), GRAY),
        ("BACKGROUND",  (0, 0), (-1, -1), LIGHT_BG),
        ("GRID",        (0, 0), (-1, -1), 0.5, BORDER),
        ("PADDING",     (0, 0), (-1, -1), 7),
        ("VALIGN",      (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(t)
    story.append(Spacer(1, 14))

    # ── Overall grade box ──
    gc = _hex(_GRADE_HEX.get(data["overall_grade"], "#374151"))
    defect_line = f"총 결함 {data['overall_defect_count']}개"
    if data["overall_max_crack_mm"]:
        defect_line += f"  ·  최대 균열폭 {data['overall_max_crack_mm']}mm"

    grade_box = [
        [Paragraph(f"종합 안전등급 : {data['overall_grade']}등급", grade_s)],
        [Paragraph(data["recommendation"], rec_s)],
        [Paragraph(defect_line, label_s)],
    ]
    gt = Table(grade_box, colWidths=[170*mm])
    gt.setStyle(TableStyle([
        ("BOX",        (0, 0), (-1, -1), 2, gc),
        ("BACKGROUND", (0, 0), (-1, -1), LIGHT_BG),
        ("PADDING",    (0, 0), (-1, -1), 12),
        ("ALIGN",      (0, 0), (-1, -1), "CENTER"),
    ]))
    story.append(gt)
    story.append(Spacer(1, 6))

    # ── Zone table ──
    story.append(Paragraph("구역별 점검 결과", h2_s))

    rows = [["구역명", "사진", "결함", "최대 균열폭", "안전등급"]]
    for z in data["zones"]:
        rows.append([
            z["name"],
            f"{z['photo_count']}장",
            f"{z['defect_count']}개",
            f"{z['max_crack_mm']}mm" if z["max_crack_mm"] else "—",
            f"{z['grade']}등급",
        ])

    zt = Table(rows, colWidths=[55*mm, 25*mm, 25*mm, 38*mm, 27*mm])
    zstyle = [
        ("FONTNAME",     (0, 0), (-1, -1), font),
        ("FONTSIZE",     (0, 0), (-1, -1), 9),
        ("BACKGROUND",   (0, 0), (-1, 0),  PRIMARY),
        ("TEXTCOLOR",    (0, 0), (-1, 0),  WHITE),
        ("FONTNAME",     (0, 0), (-1, 0),  bold),
        ("GRID",         (0, 0), (-1, -1), 0.5, BORDER),
        ("PADDING",      (0, 0), (-1, -1), 7),
        ("ALIGN",        (1, 0), (-1, -1), "CENTER"),
        ("VALIGN",       (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LIGHT_BG]),
    ]
    for i, z in enumerate(data["zones"], 1):
        zstyle.append(("TEXTCOLOR", (4, i), (4, i), _hex(_GRADE_HEX.get(z["grade"], "#374151"))))
        zstyle.append(("FONTNAME",  (4, i), (4, i), bold))
    zt.setStyle(TableStyle(zstyle))
    story.append(zt)

    # ── Summary section (standard template only) ──
    if data["template_type"] == "standard" and data["zones"]:
        story.append(Paragraph("구역 상세", h2_s))
        for z in data["zones"]:
            row_data = [
                [z["name"], f"사진 {z['photo_count']}장", f"결함 {z['defect_count']}개",
                 f"최대 균열 {z['max_crack_mm']}mm" if z["max_crack_mm"] else "균열 없음",
                 f"{z['grade']}등급"],
            ]
            dt = Table(row_data, colWidths=[42*mm, 28*mm, 28*mm, 42*mm, 30*mm])
            dt.setStyle(TableStyle([
                ("FONTNAME",   (0, 0), (-1, -1), font),
                ("FONTSIZE",   (0, 0), (-1, -1), 9),
                ("BACKGROUND", (0, 0), (-1, -1), LIGHT_BG),
                ("GRID",       (0, 0), (-1, -1), 0.5, BORDER),
                ("PADDING",    (0, 0), (-1, -1), 6),
                ("TEXTCOLOR",  (4, 0), (4, 0), _hex(_GRADE_HEX.get(z["grade"], "#374151"))),
                ("FONTNAME",   (4, 0), (4, 0), bold),
            ]))
            story.append(dt)
            story.append(Spacer(1, 4))

    # ── Footer ──
    story.append(Spacer(1, 20))
    story.append(HRFlowable(width="100%", thickness=0.5, color=BORDER, spaceAfter=6))
    footer_text = (
        f"보고서 ID: {data['report_id']}  |  CrackBot AI 자동 생성  |  "
        f"생성: {data['generated_at'].strftime('%Y-%m-%d %H:%M')}"
    )
    story.append(Paragraph(footer_text, label_s))

    doc.build(story)
    return buf.getvalue()


# ── DOCX ────────────────────────────────────────────────────────

def _generate_docx(data: dict) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    _GRADE_RGB = {
        "A": RGBColor(0x05, 0x96, 0x69),
        "B": RGBColor(0x02, 0x84, 0xC7),
        "C": RGBColor(0xD9, 0x77, 0x06),
        "D": RGBColor(0xDC, 0x26, 0x26),
        "E": RGBColor(0x7C, 0x3A, 0xED),
    }
    GRAY_RGB   = RGBColor(0x6B, 0x72, 0x80)
    PRIMARY_HX = "1D9E75"

    def _shd(cell, fill_hex: str):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill_hex)
        tcPr.append(shd)

    doc = Document()
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # Title
    h = doc.add_heading("건축물 안전점검 보고서", 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    tmpl_label = "국토안전관리원 표준 양식" if data["template_type"] == "standard" else "간이 요약 보고서"
    sub = doc.add_paragraph(tmpl_label)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in sub.runs:
        r.font.color.rgb = GRAY_RGB
        r.font.size = Pt(10)

    doc.add_paragraph()

    # Project info table
    tbl = doc.add_table(rows=3, cols=4)
    tbl.style = "Table Grid"
    rows_info = [
        ("프로젝트명", data["project_title"],    "건물명",   data["building_name"]),
        ("주소",       data["building_address"], "점검 기간", f"{data['start_date']} ~ {data['end_date']}"),
        ("점검자",     data["inspector_name"],   "생성 일시", data["generated_at"].strftime("%Y-%m-%d %H:%M")),
    ]
    for i, row_vals in enumerate(rows_info):
        row = tbl.rows[i]
        for j, val in enumerate(row_vals):
            cell = row.cells[j]
            cell.text = val
            if j % 2 == 0:
                for r in cell.paragraphs[0].runs:
                    r.font.color.rgb = GRAY_RGB
                    r.bold = True

    doc.add_paragraph()

    # Overall grade
    grade      = data["overall_grade"]
    grade_rgb  = _GRADE_RGB.get(grade, RGBColor(0, 0, 0))
    gh = doc.add_heading(f"종합 안전등급 : {grade}등급", 1)
    for r in gh.runs:
        r.font.color.rgb = grade_rgb

    rec = doc.add_paragraph(data["recommendation"])
    for r in rec.runs:
        r.bold = True

    defect_line = f"총 결함 {data['overall_defect_count']}개"
    if data["overall_max_crack_mm"]:
        defect_line += f"  ·  최대 균열폭 {data['overall_max_crack_mm']}mm"
    doc.add_paragraph(defect_line)
    doc.add_paragraph()

    # Zone table
    doc.add_heading("구역별 점검 결과", 2)
    zt = doc.add_table(rows=1 + len(data["zones"]), cols=5)
    zt.style = "Table Grid"

    headers = ["구역명", "사진 수", "결함 수", "최대 균열폭", "안전등급"]
    hdr = zt.rows[0]
    for j, h_txt in enumerate(headers):
        cell = hdr.cells[j]
        cell.text = h_txt
        _shd(cell, PRIMARY_HX)
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, zone in enumerate(data["zones"], 1):
        row = zt.rows[i]
        row.cells[0].text = zone["name"]
        row.cells[1].text = f"{zone['photo_count']}장"
        row.cells[2].text = f"{zone['defect_count']}개"
        row.cells[3].text = f"{zone['max_crack_mm']}mm" if zone["max_crack_mm"] else "—"
        row.cells[4].text = f"{zone['grade']}등급"
        zrgb = _GRADE_RGB.get(zone["grade"])
        if zrgb:
            for r in row.cells[4].paragraphs[0].runs:
                r.font.color.rgb = zrgb
                r.bold = True

    doc.add_paragraph()

    # Footer
    ft = doc.add_paragraph(
        f"보고서 ID: {data['report_id']}  |  CrackBot AI 자동 생성  |  "
        f"생성: {data['generated_at'].strftime('%Y-%m-%d %H:%M')}"
    )
    for r in ft.runs:
        r.font.color.rgb = GRAY_RGB
        r.font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Public API ──────────────────────────────────────────────────

def generate_report_files(report: Report, db: Session) -> tuple[Path | None, Path | None]:
    """
    Generate PDF and/or DOCX for the given report.
    Returns (pdf_path, docx_path) — None for types not requested.
    Files are written to the reports/ directory.
    """
    data = _gather_data(report, db)
    fmt  = str(report.format.value)  # 'pdf' | 'docx' | 'both'

    pdf_path  = None
    docx_path = None

    if fmt in ("pdf", "both"):
        pdf_bytes = _generate_pdf(data)
        pdf_path  = _REPORTS_DIR / f"{report.id}.pdf"
        pdf_path.write_bytes(pdf_bytes)

    if fmt in ("docx", "both"):
        docx_bytes = _generate_docx(data)
        docx_path  = _REPORTS_DIR / f"{report.id}.docx"
        docx_path.write_bytes(docx_bytes)

    return pdf_path, docx_path
